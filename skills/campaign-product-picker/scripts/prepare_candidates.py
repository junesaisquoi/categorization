#!/usr/bin/env python3
"""Prepare deterministic merchandising eligibility/freshness flags.

Usage:
  python prepare_candidates.py PRODUCTS.csv CATEGORIES.csv OUTPUT.csv \
      [--context-json CONTEXT.json] \
      [--destinations men-featured-22,women-featured-22] \
      [--delivery-list DELIVERY.csv]

The script collapses variant-row product exports to one row per product, detects
sale membership, detects the two most recent prior featured campaigns per
Men/Women audience from category orderHint, and marks an optional delivery item
list as the preferred candidate pool.
"""

import argparse
import csv
import json
import re
import sys
from collections import OrderedDict, defaultdict
from pathlib import Path


def die(message: str) -> None:
    print(f"ERROR: {message}", file=sys.stderr)
    raise SystemExit(2)


def truthy(value: str) -> bool:
    return str(value or "").strip().lower() in {"true", "1", "yes", "y"}


def falsy(value: str) -> bool:
    return str(value or "").strip().lower() in {"false", "0", "no", "n"}


def parse_order_hint(value: str) -> float:
    text = str(value or "").strip().replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return float("inf")


def split_categories(value: str) -> list[str]:
    return [x.strip() for x in str(value or "").split(";") if x.strip()]


def load_products(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            die("Product export has no header.")
        fields = list(reader.fieldnames)
        for required in ("key", "variants.key", "categories"):
            if required not in fields:
                die(f"Product export is missing required column: {required}")

        products = OrderedDict()
        current_key = None
        for raw in reader:
            key = (raw.get("key") or "").strip()
            if key:
                current_key = key
                products.setdefault(key, {name: (raw.get(name) or "") for name in fields})
                products[key]["key"] = key
                products[key]["_variant_signals"] = []
            elif current_key is None:
                continue
            else:
                # Fill blanks from continuation/variant rows, preserving first values.
                for name in fields:
                    value = raw.get(name) or ""
                    if not products[current_key].get(name) and value:
                        products[current_key][name] = value

            if current_key and not products[current_key].get("variants.key"):
                vkey = (raw.get("variants.key") or "").strip()
                if vkey:
                    products[current_key]["variants.key"] = vkey
            if current_key:
                products[current_key]["_variant_signals"].append({
                    "published": raw.get("attributes.published", ""),
                    "visible": raw.get("attributes.visible_flag_b2c", ""),
                    "item_out": raw.get("attributes.item_out", ""),
                })

    return fields, products


def load_categories(path: Path) -> dict[str, dict]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            die("Category export has no header.")
        for required in ("key", "parent.key", "orderHint"):
            if required not in reader.fieldnames:
                die(f"Category export is missing required column: {required}")
        return {(r.get("key") or "").strip(): r for r in reader if (r.get("key") or "").strip()}


# A discounted product is not always in a category with the word "sale" in it. Friends &
# Family and app early access are price-reduced windows on the same footing, so they
# belong in the same exclusion. Keep this identical to SALE_MARKERS in
# ../../category-builder/scripts/rules.py, or the two skills will disagree about which
# products are on sale.
SALE_MARKERS = ("sale", "friends-family", "early-access")
_MARKER_ALTERNATION = "|".join(SALE_MARKERS)


def sale_category_set(categories: dict[str, dict]) -> set[str]:
    """Return keys that are sale categories or descendants of sale categories."""
    sale_re = re.compile(rf"(^|-)({_MARKER_ALTERNATION})($|-)", re.I)
    sale_name_re = re.compile(rf"\b({_MARKER_ALTERNATION.replace('-', '[ -]?')})\b", re.I)
    sale = set()
    for key, row in categories.items():
        localized_names = [value for field, value in row.items() if field == "name" or field.startswith("name.")]
        if sale_re.search(key) or any(sale_name_re.search(str(value or "")) for value in localized_names):
            sale.add(key)
    changed = True
    while changed:
        changed = False
        for key, row in categories.items():
            parent = (row.get("parent.key") or "").strip()
            if parent in sale and key not in sale:
                sale.add(key)
                changed = True
    return sale


def category_is_at_or_below(categories: dict[str, dict], key: str, root: str) -> bool:
    """Return whether key is root or a descendant of root, guarding against cycles."""
    current = key
    seen = set()
    while current and current not in seen:
        if current == root:
            return True
        seen.add(current)
        current = (categories.get(current, {}).get("parent.key") or "").strip()
    return False


def regular_audience_categories(
    assigned: set[str], categories: dict[str, dict], sale_categories: set[str], audience: str
) -> list[str]:
    """Return assigned regular storefront categories proving audience eligibility.

    Featured-history and Sale categories never establish regular audience eligibility.
    """
    audience_root = audience
    featured_root = f"{audience}-featured"
    return sorted(
        key
        for key in assigned
        if key not in sale_categories
        and category_is_at_or_below(categories, key, audience_root)
        and not category_is_at_or_below(categories, key, featured_root)
    )


def referenced_category_counts(products) -> dict[str, int]:
    members = defaultdict(set)
    for key, row in products.items():
        for cat in split_categories(row.get("categories", "")):
            members[cat].add(key)
    return {cat: len(keys) for cat, keys in members.items()}


def is_stale_inactive_label(name: str) -> bool:
    text = str(name or "").lower()
    return "(inactive)" in text or "(inaktiv)" in text


def recent_featured(categories, counts, audience: str, destinations: set[str], limit=2):
    parent = f"{audience}-featured"
    candidates = []
    for key, row in categories.items():
        if (row.get("parent.key") or "").strip() != parent:
            continue
        if key in destinations:
            continue
        if counts.get(key, 0) <= 0:  # "run" campaigns must actually contain products
            continue
        active = (row.get("custom.fields.active") or "").strip()
        if active and not truthy(active):
            continue
        if is_stale_inactive_label(row.get("name.en", "")):
            continue
        hint = parse_order_hint(row.get("orderHint", ""))
        if hint == float("inf"):
            continue
        candidates.append({
            "key": key,
            "name": (row.get("name.en") or key).strip(),
            "orderHint": row.get("orderHint", ""),
            "productCount": counts.get(key, 0),
        })
    candidates.sort(key=lambda x: (parse_order_hint(x["orderHint"]), x["key"]))
    return candidates[:limit]


def featured_memberships(products, audience: str) -> dict[str, list[str]]:
    prefix = f"{audience}-featured-"
    result = {}
    for key, row in products.items():
        cats = [c for c in split_categories(row.get("categories", "")) if c.startswith(prefix)]
        result[key] = cats
    return result


def load_delivery_keys(path: Path, known_keys: set[str]) -> set[str]:
    if not path:
        return set()
    if not path.exists():
        die(f"Delivery list not found: {path}")

    suffix = path.suffix.lower()
    found = set()

    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8-sig", errors="ignore")
        for key in known_keys:
            if key in text:
                found.add(key)
        return found

    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except Exception:
            die("XLSX delivery lists require openpyxl; convert the list to CSV if unavailable.")
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    text = str(cell or "").strip()
                    if text in known_keys:
                        found.add(text)
        return found

    # CSV/TSV: scan all cells so the source column name does not matter.
    delimiter = "\t" if suffix == ".tsv" else ","
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f, delimiter=delimiter)
        for row in reader:
            for cell in row:
                text = str(cell or "").strip()
                if text in known_keys:
                    found.add(text)
    return found


def current_online(row: dict) -> bool:
    """Conservative current-online test using fields available in CT exports."""
    published = row.get("published", "")
    visible = row.get("attributes.main_visible_flag_b2c", "")
    item_out = row.get("attributes.main_item_out", "")

    if published and not truthy(published):
        return False
    if visible and not truthy(visible):
        return False
    if item_out and truthy(item_out):
        return False

    variant_signals = row.get("_variant_signals") or []
    variants_with_status = []
    for signals in variant_signals:
        values = [signals.get("published", ""), signals.get("visible", ""), signals.get("item_out", "")]
        if not any(str(value or "").strip() for value in values):
            continue
        available = True
        if signals.get("published") and not truthy(signals["published"]):
            available = False
        if signals.get("visible") and not truthy(signals["visible"]):
            available = False
        if signals.get("item_out") and truthy(signals["item_out"]):
            available = False
        variants_with_status.append(available)
    if variants_with_status and not any(variants_with_status):
        return False
    return True


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("products")
    ap.add_argument("categories")
    ap.add_argument("output_csv")
    ap.add_argument("--context-json")
    ap.add_argument("--destinations", default="")
    ap.add_argument("--delivery-list")
    args = ap.parse_args()

    product_path = Path(args.products)
    category_path = Path(args.categories)
    if not product_path.exists():
        die(f"Product export not found: {product_path}")
    if not category_path.exists():
        die(f"Category export not found: {category_path}")

    fields, products = load_products(product_path)
    categories = load_categories(category_path)
    counts = referenced_category_counts(products)
    destinations = {x.strip() for x in args.destinations.split(",") if x.strip()}
    sale_cats = sale_category_set(categories)
    delivery_keys = load_delivery_keys(Path(args.delivery_list), set(products)) if args.delivery_list else set()

    recent = {
        aud: recent_featured(categories, counts, aud, destinations, 2)
        for aud in ("men", "women")
    }
    recent_cat_keys = {aud: {x["key"] for x in vals} for aud, vals in recent.items()}
    featured = {aud: featured_memberships(products, aud) for aud in ("men", "women")}

    helper_fields = [
        "_online_now",
        "_delivery_priority",
        "_sale_hard_no",
        "_men_audience_eligible",
        "_women_audience_eligible",
        "_men_regular_categories",
        "_women_regular_categories",
        "_men_recent_two_hard_no",
        "_women_recent_two_hard_no",
        "_men_featured_history",
        "_women_featured_history",
    ]
    out_fields = fields + [h for h in helper_fields if h not in fields]

    output_path = Path(args.output_csv)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    sale_keys = set()
    recent_product_keys = {"men": set(), "women": set()}
    online_keys = set()
    regular_category_evidence = {"men": {}, "women": {}}

    with output_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=out_fields, extrasaction="ignore")
        writer.writeheader()
        for key, row in products.items():
            cats = set(split_categories(row.get("categories", "")))
            sale = bool(cats & sale_cats) or any(
                re.search(rf"(^|-)({_MARKER_ALTERNATION})($|-)", c, re.I) for c in cats)
            if sale:
                sale_keys.add(key)
            online = current_online(row)
            if online:
                online_keys.add(key)

            flags = dict(row)
            flags["_online_now"] = "true" if online else "false"
            flags["_delivery_priority"] = "true" if key in delivery_keys else "false"
            flags["_sale_hard_no"] = "true" if sale else "false"
            for aud in ("men", "women"):
                regular = regular_audience_categories(cats, categories, sale_cats, aud)
                if regular:
                    regular_category_evidence[aud][key] = regular
                flags[f"_{aud}_audience_eligible"] = "true" if regular else "false"
                flags[f"_{aud}_regular_categories"] = ";".join(regular)
                hit = bool(cats & recent_cat_keys[aud])
                if hit:
                    recent_product_keys[aud].add(key)
                flags[f"_{aud}_recent_two_hard_no"] = "true" if hit else "false"
                flags[f"_{aud}_featured_history"] = ";".join(featured[aud].get(key, []))
            writer.writerow(flags)

    context = {
        "destinations": sorted(destinations),
        "recent_featured": recent,
        "hard_exclusions": {
            "sale_keys": sorted(sale_keys),
            "recent_two_keys": {aud: sorted(keys) for aud, keys in recent_product_keys.items()},
        },
        "priority": {
            "delivery_list_keys": sorted(delivery_keys),
        },
        "eligibility": {
            "online_now_keys": sorted(online_keys),
            "regular_audience_categories": regular_category_evidence,
        },
    }
    context_path = Path(args.context_json) if args.context_json else output_path.with_suffix(".context.json")
    context_path.write_text(json.dumps(context, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"Prepared {len(products)} product-level rows: {output_path}")
    print(f"Context: {context_path}")
    print(f"Sale hard-no products: {len(sale_keys)}")
    if delivery_keys:
        print(f"Delivery priority products matched: {len(delivery_keys)}")
    for aud in ("men", "women"):
        print(f"{aud.title()} regular-audience products: {len(regular_category_evidence[aud])}")
    for aud in ("men", "women"):
        labels = ", ".join(f"{x['key']} ({x['name']}, {x['orderHint']})" for x in recent[aud]) or "none"
        print(f"{aud.title()} recent two campaigns: {labels}")
        print(f"{aud.title()} products in recent two: {len(recent_product_keys[aud])}")


if __name__ == "__main__":
    main()
